"""
Offline tests for filling Tennents' own pricer from the master.

Covers: the off-invoice lands in column J of the right "BDM Copy" rows (matched by
code, a compound code, a leading-zero code and, when unique, by name); blank for
a product not sold; managed pubs take the whole discount; pub name and date are
written; Tennents' WSP and total are left as sent; the check lists a product the
pub sells that the pricer has no row for, a total that differs from the master's
and an off-invoice above Tennents' total; the file comes back locked with the
password and opens with it; the .zip carries one file per pub. Uses the real
template in pricer_template/. No Airtable network access.

    python test_tennents_pricer.py
"""
import io
import sys
import zipfile
from datetime import date

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import msoffcrypto
from openpyxl import load_workbook

import tennents_pricer as tp
from tennents_master import SiteInfo, SitePrice, SkuRate, TennentsMaster

PASS = True


def _check(label, cond, detail=""):
    global PASS
    PASS &= bool(cond)
    print(f"  [{'ok' if cond else 'FAIL'}] {label}{(' — ' + detail) if detail and not cond else ''}")


def make_master() -> TennentsMaster:
    skus = [
        # "090425" in the master, 90425 on the pricer: leading-zero drift
        SkuRate("090425", "", "Tennent's", "Tennent's Lager", "11G / 22G",
                None, 4.0, 641.62, 302.00, True, "C&C", 12.33, 314.33),
        SkuRate("400889", "", "Tennent's", "Jeffreys Lager", "50L",
                0.3055, 3.4, 538.39, 291.52, True, "C&C", 11.45, 302.97),
        # pricer row carries "400557" alone on BDM Copy; master holds alt 400751
        SkuRate("400557", "400751", "Menabrea", "Menabrea", "50L",
                0.3055, 4.8, 805.6, 332.0, True, "3rd party", 0.0, 332.0),
        # code differs from the pricer's ("GUI002") only via the alt: name match not needed
        SkuRate("GUI002", "GUI003", "Guinness", "Guinness Keg", "50L / 30L",
                None, 4.1, 779.24, 250.00, True, "3rd party", 0.0, 250.00),
        # on no pricer row at all
        SkuRate("401300", "", "Kingfisher", "Kingfisher", "30L",
                0.1833, 4.8, 900.0, 300.0, True, "3rd party", 0.0, 300.0),
        # pricer code "T00045238" unknown to the master; matched by unique name
        SkuRate("CARL01", "", "Carlsberg", "Carlsberg Pilsner", "50L",
                0.3055, 3.4, 796.51, 410.0, True, "3rd party", 0.0, 410.0),
    ]
    sites = [
        SiteInfo("17599415", "MALLROAD HOUSE", "Tenanted (TBC)", "Standard split"),
        SiteInfo("11110002", "MANAGED HOUSE", "MANAGED (confirmed)", "ALL OFF-INVOICE, zero retro"),
    ]
    site_prices = [
        SitePrice("17599415", "MALLROAD HOUSE", "090425", "T.Lager", 144.69),
        SitePrice("17599415", "MALLROAD HOUSE", "400889", "Jeffreys", 118.0),
        SitePrice("17599415", "MALLROAD HOUSE", "400751", "Menabrea", 200.0),   # via alt code
        SitePrice("17599415", "MALLROAD HOUSE", "GUI002", "Guinness", 0.0),     # £0 = not sold
        SitePrice("17599415", "MALLROAD HOUSE", "401300", "Kingfisher", 120.0),
        SitePrice("17599415", "MALLROAD HOUSE", "CARL01", "Carlsberg", 450.0),  # above total
    ]
    return TennentsMaster("vTEST", "fixture", skus, sites, [], site_prices)


def bdm(data: bytes, password: str | None = None):
    if password:
        f = msoffcrypto.OfficeFile(io.BytesIO(data))
        f.load_key(password=password)
        out = io.BytesIO()
        f.decrypt(out)
        data = out.getvalue()
    return load_workbook(io.BytesIO(data))["BDM Copy"]


def row_of(ws, brand):
    for r in range(tp.FIRST_ROW, ws.max_row + 1):
        if str(ws.cell(r, tp.COL_BRAND).value or "").strip() == brand:
            return r
    raise AssertionError(brand)


def run():
    m = make_master()
    mall = m.site_for_account("17599415")
    data, check = tp.build_pricer_bytes(m, mall, date(2026, 10, 12), password="")
    ws = bdm(data)

    _check("pub name in J2", ws["J2"].value == "MALLROAD HOUSE", repr(ws["J2"].value))
    _check("date in L2", ws["L2"].value is not None and ws["L2"].value.date() == date(2026, 10, 12))
    _check("Tennents Lager off 144.69 (leading-zero code)", ws.cell(row_of(ws, "Tennents Lager"), 10).value == 144.69)
    _check("Jeffreys off 118", ws.cell(row_of(ws, "Jeffreys Lager"), 10).value == 118.0)
    _check("Menabrea off via alt code", ws.cell(row_of(ws, "Menabrea"), 10).value == 200.0)
    _check("Guinness at £0 left blank (not sold)", ws.cell(row_of(ws, "Guinness Keg"), 10).value is None)
    _check("Carlsberg matched by unique name", ws.cell(row_of(ws, "Carlsberg Pilsner"), 10).value == 450.0)
    _check("unpriced product left blank", ws.cell(row_of(ws, "Heverlee"), 10).value is None)
    r = row_of(ws, "Tennents Lager")
    _check("Tennents' WSP untouched", ws.cell(r, 5).value == 641.62)
    _check("Tennents' total untouched", ws.cell(r, 6).value == 314.33)
    _check("retro formula kept", ws.cell(r, 11).value == f"=F{r}-J{r}")

    names = [s.product for s, _ in check.not_on_pricer]
    _check("Kingfisher listed as not on the pricer", names == ["Kingfisher"], repr(names))
    differs = [ln.brand for ln in check.total_differs]
    _check("Guinness not flagged (not sold)", "Guinness Keg" not in differs, repr(differs))
    _check("Jeffreys total equal: no flag", "Jeffreys Lager" not in differs, repr(differs))
    above = [ln.brand for ln in check.off_above_total]
    _check("Carlsberg off above Tennents' total flagged", above == ["Carlsberg Pilsner"], repr(above))
    _check("unknown pricer rows kept as unmatched", any(ln.brand == "Heverlee" for ln in check.unmatched_rows))

    # check_pricer (the page) agrees with what the download wrote
    c2 = tp.check_pricer(m, mall)
    _check("check matches the build", [ln.off for ln in c2.lines] == [ln.off for ln in check.lines])

    # managed: every matched product carries the whole discount, nothing "not on pricer"
    man = m.site_for_account("11110002")
    data_m, check_m = tp.build_pricer_bytes(m, man, password="")
    wm = bdm(data_m)
    _check("managed: Guinness takes the whole 250", wm.cell(row_of(wm, "Guinness Keg"), 10).value == 250.0)
    _check("managed: no not-on-pricer list", check_m.not_on_pricer == [])

    # locked with the password, as Tennents send it
    locked, _ = tp.build_pricer_bytes(m, mall)
    _check("default file is encrypted", msoffcrypto.OfficeFile(io.BytesIO(locked)).is_encrypted())
    _check("opens with the Tennents password", bdm(locked, tp.DEFAULT_PASSWORD)["J2"].value == "MALLROAD HOUSE")

    z = zipfile.ZipFile(io.BytesIO(tp.build_pricer_zip_bytes(m, m.sites, date(2026, 10, 12), password="")))
    _check("zip: one pricer per pub", sorted(z.namelist()) == [
        "FB TAVERNS - PRICER - MALLROAD HOUSE 2026-10-12.xlsx",
        "FB TAVERNS - PRICER - MANAGED HOUSE 2026-10-12.xlsx"], repr(z.namelist()))


if __name__ == "__main__":
    run()
    print("\nALL PASS" if PASS else "\nSOME FAILED")
    sys.exit(0 if PASS else 1)
